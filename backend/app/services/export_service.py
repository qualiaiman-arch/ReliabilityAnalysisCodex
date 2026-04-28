from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from pptx import Presentation


def export_excel(calculation: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Reliability Results"

    ws.append(["System", calculation["system_name"]])
    ws.append(["R", calculation["system_r"]])
    ws.append(["Q", calculation["system_q"]])
    ws.append([])
    ws.append(["Assembly", "R", "Q", "Unintended Op", "Redundancy", "Assumptions/Notes"])

    for a in calculation["assemblies"]:
        ws.append([
            a["name"],
            a["r"],
            a["q"],
            a["unintended_operation_probability"],
            a["redundancy_type"],
            "MVP constant failure rate assumption",
        ])

    ws.append([])
    ws.append(["Recommendations"])
    for r in calculation["recommendations"]:
        ws.append([r])

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def export_powerpoint(system_payload: dict, calculation: dict) -> bytes:
    prs = Presentation()

    def add_bullet_slide(title: str, bullets: list[str]):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        tf = slide.placeholders[1].text_frame
        tf.clear()
        for i, b in enumerate(bullets):
            if i == 0:
                tf.text = b
            else:
                p = tf.add_paragraph()
                p.text = b

    add_bullet_slide("System Description", [
        f"Name: {system_payload['name']}",
        f"Mission: {system_payload.get('mission', '')}",
        f"Environment: {system_payload.get('operating_environment', '')}",
    ])
    add_bullet_slide("Assumptions", [
        "Exponential reliability model with constant failure rate",
        "No common-cause failure",
        "No loops in RBD",
    ])
    add_bullet_slide("Duty Cycle", [
        "Assembly-specific active/passive/rest times used.",
        "Risk contributions include active, passive, and rest exposure.",
    ])
    add_bullet_slide("Assemblies and Components", [a["name"] for a in calculation["assemblies"]])
    add_bullet_slide("RBD", ["RBD structure is user-defined and acyclic."])
    add_bullet_slide("FTA - Failure to Operate", ["Auto-derived from RBD (series->OR, parallel 1oo2->AND)."])
    add_bullet_slide("Results Table", [f"System R={calculation['system_r']:.6f}, Q={calculation['system_q']:.6f}"])
    add_bullet_slide("Sensitivity Ranking", [f"{a['name']}: Q={a['q']:.6f}" for a in calculation["assemblies"]])
    add_bullet_slide("Design Recommendations", calculation["recommendations"] or ["No recommendation generated."])
    add_bullet_slide("Test Recommendations", [
        "Focus verification on dominant Q contributors.",
        "Perform acceptance tests for single points of failure.",
    ])

    out = BytesIO()
    prs.save(out)
    return out.getvalue()
